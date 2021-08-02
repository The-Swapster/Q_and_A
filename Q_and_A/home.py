from flask import Flask, render_template, request, send_file
import PyPDF2
import docx
#import nltk
#nltk.download('stopwords')
from Questgen import main
from sense2vec import Sense2Vec
from fpdf import FPDF
import os

s2v = Sense2Vec().from_disk("C:/Users/Swapnil/Desktop/work/s2v_old")
app = Flask(__name__)
app.config['UPLOAD_EXTEBSIONS'] = ['.pdf', '.docx']
app.config['UPLOAD_FOLDER'] = r"C:\Users\Swapnil\Desktop\work"
@app.route("/")
def homepage():
  return render_template("home.html")

@app.route("/", methods=['POST'])
def upload_file():
    uploaded_file = request.files['myFile']
    name = uploaded_file.filename
    qtype = request.form['type']
    text = request.form['text']
    c = request.form['count']
    c = int(c)
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", size=15)
    if text=="":
        if name[name.rfind("."):] == ".pdf":
            pdfReader = PyPDF2.PdfFileReader(uploaded_file)
            pageObj = pdfReader.getPage(0)
            # extracting text from page
            page_content = pageObj.extractText()
        elif name[name.rfind("."):] == ".docx":
            doc = docx.Document(uploaded_file)
            fullText = []
            for para in doc.paragraphs:
                fullText.append(para.text)
            page_content = '\n'.join(fullText)
        else:
            msg = "Incorrect File Type"
    else:
        page_content = text

    if qtype == 'boolean':
        qb = main.BoolQGen()
        payload = {
            "input_text" : page_content,
            "max_questions" : c
        }
        output = qb.predict_boolq(payload)
        answer = main.AnswerPredictor()
        ans = []
        msg=output['Boolean Questions']
        for i in output['Boolean Questions']:
            payload = {
                "input_text":page_content,
                "input_question":i
            }
            ans.append(answer.predict_answer(payload))
        count=2
        pdf.cell(200,10,txt="Boolean Questions", ln=1, align='C')
        counter = 2
        for i in range(len(msg)):
            pdf.cell(200,10,txt=str(i+1), ln=counter, align='L')
            counter = counter+1
            pdf.cell(200,10,txt='Question ', ln=counter, align='L')
            counter = counter+1
            pdf.cell(200,10,txt=msg[i], ln=counter, align='L')
            counter = counter+1
            pdf.cell(200,10,txt='Answer ', ln=counter, align='L')
            counter = counter+1
            pdf.cell(200,10,txt=ans[i], ln=counter, align='L')
            counter = counter+1
        pdf.output('questions.pdf')
        return render_template("home.html", msg=msg, ans=ans, count=count)
    elif qtype == 'mcq':
        mc = main.QGen()
        payload = {
            "input_text" : page_content,
            "max_questions" : c
        }
        output = mc.predict_mcq(payload)
        op = []
        ans = []
        msg = []
        for i in range(len(output['questions'])):
            msg.append(output['questions'][i]['question_statement'])
            a = output['questions'][i]['extra_options']
            a.append(output['questions'][i]['answer'])
            op.append(a)
            ans.append(output['questions'][i]['answer'])
        count=3
        pdf.cell(200,10,txt="Multiple Choice Questions", ln=1, align='C')
        counter = 2
        for i in range(len(msg)):
            pdf.cell(200,10,txt=str(i+1), ln=counter, align='L')
            counter = counter+1
            pdf.cell(200,10,txt='Question ', ln=counter, align='L')
            counter = counter+1
            pdf.cell(200,10,txt=msg[i], ln=counter, align='L')
            counter = counter+1
            pdf.cell(200,10,txt='Options ', ln=counter, align='L')
            counter = counter+1
            pdf.cell(200,10,txt=str(op[i]), ln=counter, align='L')
            counter = counter+1
            pdf.cell(200,10,txt='Answer ', ln=counter, align='L')
            counter = counter+1
            pdf.cell(200,10,txt=ans[i], ln=counter, align='L')
            counter = counter+1
        pdf.output('questions.pdf')
        return render_template("home.html", op=op, msg=msg, ans=ans, count=count)
    elif qtype == 'faq':
        fc = main.QGen()
        payload = {
            "input_text" : page_content,
            "max_questions" : c
        }
        output = fc.predict_shortq(payload)
        msg = []
        ans = []
        for i in range(len(output['questions'])):
            msg.append(output['questions'][i]['Question'])
            ans.append(output['questions'][i]['Answer'])
        count=2
        pdf.cell(200,10,txt="FAQs", ln=1, align='C')
        counter = 2
        for i in range(len(msg)):
            pdf.cell(200,10,txt=str(i+1), ln=counter, align='L')
            counter = counter+1
            pdf.cell(200,10,txt='Question ', ln=counter, align='L')
            counter = counter+1
            pdf.cell(200,10,txt=msg[i], ln=counter, align='L')
            counter = counter+1
            pdf.cell(200,10,txt='Answer ', ln=counter, align='L')
            counter = counter+1
            pdf.cell(200,10,txt=ans[i], ln=counter, align='L')
            counter = counter+1
        pdf.output('questions.pdf')
        return render_template("home.html", msg=msg, ans=ans, count=count)

@app.route('/download', methods=['GET'])
def download():
    path = os.path.join(r"C:\Users\Swapnil\Desktop\work", "questions.pdf")
    return send_file(path, as_attachment=True)

if __name__ == "__main__":
  app.run(debug=True)